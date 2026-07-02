"""엑셀(.xlsx) 저장."""
from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from ..schema import COLUMNS, Product

# 사람이 보기 좋은 한글 헤더 (COLUMNS 와 1:1 대응)
_HEADER_KO = {
    "source": "소스",
    "product_code": "상품코드",
    "keyword": "검색어",
    "name": "상품명",
    "price": "대표판매가",
    "list_price": "정상가",
    "sale_price": "할인가",
    "coupon_price": "쿠폰가",
    "shipping_fee": "배송비",
    "sourcing_cost": "소싱원가(입력)",
    "estimated_margin": "추정마진(계산)",
    "image_url": "이미지URL(참고용)",
    "product_url": "상품URL",
    "seller": "판매자",
    "review_count": "리뷰수",
    "rating": "평점",
    "category": "카테고리",
    "rank": "순위",
    "delivery": "배송",
    "collected_at": "수집시각",
}


def save_xlsx(products: Sequence[Product], path: str | Path) -> Path:
    """상품 목록을 xlsx로 저장하고 경로 반환."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "상품목록"

    # 헤더
    headers = [_HEADER_KO.get(c, c) for c in COLUMNS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    # 데이터
    for p in products:
        row = p.to_row()
        ws.append([row[c] for c in COLUMNS])

    # 컬럼 폭 대략 조정
    widths = {"name": 45, "product_url": 30, "image_url": 30, "seller": 18}
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 12)

    wb.save(str(path))
    return path
